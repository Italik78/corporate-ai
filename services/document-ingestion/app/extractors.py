"""Format-specific extractors that emit the normalized document blocks.

The ingestion pipeline deliberately converts every supported format into the same
NormalizedBlock structure before chunking. Extractors never execute document
content and preserve source-local provenance in each block.
"""

from __future__ import annotations

import csv
import io
import re
from pathlib import Path
from typing import Any

from .models import NormalizedBlock


TEXT_SUFFIXES = {".txt", ".md", ".markdown"}
SUPPORTED_SUFFIXES = TEXT_SUFFIXES | {".docx", ".xlsx", ".pptx", ".csv"}


def extract_text(filename: str, data: bytes) -> list[NormalizedBlock]:
    suffix = Path(filename).suffix.lower()
    if suffix not in TEXT_SUFFIXES:
        raise ValueError("UNSUPPORTED_EXTRACTOR")

    return _extract_plain_text(filename, data)


def extract_document(filename: str, data: bytes) -> list[NormalizedBlock]:
    """Route a validated file to its format-specific extractor."""
    suffix = Path(filename).suffix.lower()

    if suffix in TEXT_SUFFIXES:
        return _extract_plain_text(filename, data)
    if suffix == ".csv":
        return _extract_csv(filename, data)
    if suffix == ".docx":
        return _extract_docx(filename, data)
    if suffix == ".xlsx":
        return _extract_xlsx(filename, data)
    if suffix == ".pptx":
        return _extract_pptx(filename, data)

    raise ValueError("UNSUPPORTED_FILE_TYPE")


def _decode_text(data: bytes) -> str:
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("utf-8-sig", errors="replace")


def _extract_plain_text(filename: str, data: bytes) -> list[NormalizedBlock]:
    suffix = Path(filename).suffix.lower()
    text = _decode_text(data)
    if not text.strip():
        return []

    blocks: list[NormalizedBlock] = []
    section: str | None = None
    counter = 0

    for line_no, raw in enumerate(text.splitlines(), start=1):
        line = raw.strip()
        if not line:
            continue

        if suffix in {".md", ".markdown"} and line.startswith("#"):
            section = line.lstrip("#").strip()
            counter += 1
            blocks.append(
                NormalizedBlock(
                    block_id=f"b-{counter:06d}",
                    block_type="heading",
                    content=section,
                    section=section,
                    provenance={
                        "line": line_no,
                        "source_format": "markdown",
                    },
                )
            )
            continue

        counter += 1
        blocks.append(
            NormalizedBlock(
                block_id=f"b-{counter:06d}",
                content=line,
                section=section,
                provenance={
                    "line": line_no,
                    "source_format": "text",
                },
            )
        )

    return blocks


def _clean_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _table_content(headers: list[str], rows: list[list[str]]) -> str:
    """Serialize a table deterministically while retaining column context."""
    if headers:
        header_line = " | ".join(headers)
        separator = " | ".join("---" for _ in headers)
        body = [" | ".join(row) for row in rows]
        return "\n".join([header_line, separator, *body]).strip()

    return "\n".join(" | ".join(row) for row in rows).strip()


def _extract_csv(filename: str, data: bytes) -> list[NormalizedBlock]:
    text = _decode_text(data)
    if not text.strip():
        return []

    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel

    reader = csv.reader(io.StringIO(text), dialect)
    rows = [[_clean_cell(cell) for cell in row] for row in reader]
    rows = [row for row in rows if any(cell for cell in row)]
    if not rows:
        return []

    headers = rows[0]
    body = rows[1:]
    width = max(len(row) for row in rows)
    headers = headers + [""] * (width - len(headers))

    normalized_rows: list[list[str]] = []
    for row in body:
        normalized_rows.append(row + [""] * (width - len(row)))

    content = _table_content(headers, normalized_rows)
    return [
        NormalizedBlock(
            block_id="csv-table-000001",
            block_type="table",
            content=content,
            section=Path(filename).stem,
            provenance={
                "source_format": "csv",
                "row_count": len(normalized_rows),
                "column_count": width,
                "header": headers,
            },
        )
    ]


def _extract_docx(filename: str, data: bytes) -> list[NormalizedBlock]:
    from docx import Document
    from docx.document import Document as DocumentType
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P

    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise ValueError("DOCX_PARSE_ERROR") from exc

    def iter_body_blocks(parent: DocumentType):
        body = parent.element.body
        for child in body.iterchildren():
            if isinstance(child, CT_P):
                yield Paragraph(child, parent)
            elif isinstance(child, CT_Tbl):
                yield Table(child, parent)

    blocks: list[NormalizedBlock] = []
    section: str | None = None
    paragraph_index = 0
    table_index = 0
    block_index = 0

    for item in iter_body_blocks(document):
        if isinstance(item, Paragraph):
            paragraph_index += 1
            text = item.text.strip()
            if not text:
                continue

            block_index += 1
            style_name = item.style.name or ""
            style = style_name.lower()
            is_heading = "heading" in style or style.startswith("title")
            if is_heading:
                section = text

            blocks.append(
                NormalizedBlock(
                    block_id=f"docx-p-{paragraph_index:06d}",
                    block_type="heading" if is_heading else "text",
                    content=text,
                    section=section,
                    provenance={
                        "source_format": "docx",
                        "paragraph_index": paragraph_index,
                        "style": style_name,
                    },
                )
            )
            continue

        table_index += 1
        rows = [
            [_clean_cell(cell.text.replace("\n", " ")) for cell in row.cells]
            for row in item.rows
        ]
        rows = [row for row in rows if any(cell for cell in row)]
        if not rows:
            continue

        width = max(len(row) for row in rows)
        headers = rows[0] + [""] * (width - len(rows[0]))
        normalized_rows = [row + [""] * (width - len(row)) for row in rows[1:]]
        block_index += 1
        blocks.append(
            NormalizedBlock(
                block_id=f"docx-table-{table_index:06d}",
                block_type="table",
                content=_table_content(headers, normalized_rows),
                section=section,
                provenance={
                    "source_format": "docx",
                    "table_index": table_index,
                    "row_count": len(normalized_rows),
                    "column_count": width,
                    "header": headers,
                },
            )
        )

    return blocks


def _extract_xlsx(filename: str, data: bytes) -> list[NormalizedBlock]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(
            io.BytesIO(data),
            read_only=True,
            data_only=False,
        )
    except Exception as exc:
        raise ValueError("XLSX_PARSE_ERROR") from exc

    blocks: list[NormalizedBlock] = []
    for sheet_index, worksheet in enumerate(workbook.worksheets, start=1):
        rows: list[list[str]] = []
        for row in worksheet.iter_rows(values_only=True):
            values = [_clean_cell(value) for value in row]
            if any(values):
                rows.append(values)

        if not rows:
            continue

        width = max(len(row) for row in rows)
        headers = rows[0] + [""] * (width - len(rows[0]))
        normalized_rows = [row + [""] * (width - len(row)) for row in rows[1:]]

        blocks.append(
            NormalizedBlock(
                block_id=f"xlsx-{sheet_index:04d}-table-000001",
                block_type="table",
                content=_table_content(headers, normalized_rows),
                section=worksheet.title,
                provenance={
                    "source_format": "xlsx",
                    "sheet": worksheet.title,
                    "sheet_index": sheet_index,
                    "row_count": len(normalized_rows),
                    "column_count": width,
                    "header": headers,
                },
            )
        )

    workbook.close()
    return blocks


def _extract_pptx(filename: str, data: bytes) -> list[NormalizedBlock]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    try:
        presentation = Presentation(io.BytesIO(data))
    except Exception as exc:
        raise ValueError("PPTX_PARSE_ERROR") from exc

    blocks: list[NormalizedBlock] = []

    for slide_index, slide in enumerate(presentation.slides, start=1):
        title = slide.shapes.title.text.strip() if slide.shapes.title else None
        section = title or f"Slide {slide_index}"
        block_index = 0

        for shape_index, shape in enumerate(slide.shapes, start=1):
            shape_type = getattr(shape, "shape_type", None)

            if getattr(shape, "has_text_frame", False):
                text = shape.text.strip()
                if text:
                    block_index += 1
                    blocks.append(
                        NormalizedBlock(
                            block_id=f"pptx-s{slide_index:04d}-b{block_index:04d}",
                            block_type="heading" if shape is slide.shapes.title else "text",
                            content=text,
                            page=slide_index,
                            section=section,
                            provenance={
                                "source_format": "pptx",
                                "slide": slide_index,
                                "shape": shape_index,
                                "shape_type": str(shape_type),
                            },
                        )
                    )

            if shape_type == MSO_SHAPE_TYPE.TABLE and getattr(shape, "has_table", False):
                rows = []
                for row in shape.table.rows:
                    rows.append([_clean_cell(cell.text) for cell in row.cells])
                rows = [row for row in rows if any(cell for cell in row)]
                if not rows:
                    continue

                width = max(len(row) for row in rows)
                headers = rows[0] + [""] * (width - len(rows[0]))
                normalized_rows = [
                    row + [""] * (width - len(row)) for row in rows[1:]
                ]
                block_index += 1
                blocks.append(
                    NormalizedBlock(
                        block_id=f"pptx-s{slide_index:04d}-table{block_index:04d}",
                        block_type="table",
                        content=_table_content(headers, normalized_rows),
                        page=slide_index,
                        section=section,
                        provenance={
                            "source_format": "pptx",
                            "slide": slide_index,
                            "shape": shape_index,
                            "row_count": len(normalized_rows),
                            "column_count": width,
                            "header": headers,
                        },
                    )
                )

    return blocks
