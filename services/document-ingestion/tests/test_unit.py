import asyncio
import csv
import io

import pytest

from app.security import sha256_bytes, validate_filename, validate_extension
from app.extractors import extract_document, extract_text
from app.chunker import chunk_document
from app.models import DocumentMetadata, LifecycleStatus, NormalizedDocument
from app import pipeline


def test_hash_deterministic():
    assert sha256_bytes(b"abc") == sha256_bytes(b"abc")
    assert sha256_bytes(b"abc") != sha256_bytes(b"abcd")


def test_filename_and_extension():
    assert validate_filename("../test.md") == "test.md"
    assert validate_extension("test.md", {".md"}) == ".md"


def test_markdown_extraction():
    blocks = extract_text("# Заглавие\n\nТекст за документа.".encode("utf-8"))
    assert len(blocks) == 2
    assert blocks[0].block_type == "heading"
    assert blocks[1].section == "Заглавие"


def _document(text: bytes) -> NormalizedDocument:
    blocks = extract_text("test.txt", text)
    meta = DocumentMetadata(
        document_id="d1",
        title="t",
        created_at="now",
        updated_at="now",
    )
    return NormalizedDocument(
        document_id="d1",
        source_file="test.txt",
        media_type="text/plain",
        content_hash="x",
        metadata=meta,
        blocks=blocks,
    )


def test_chunking():
    chunks = chunk_document(_document(("думата " * 400).encode()), target_words=100, overlap_words=10)
    assert len(chunks) >= 4
    assert all(c.document_id == "d1" for c in chunks)
    assert all(len(c.content.split()) <= 100 for c in chunks)


def test_chunking_preserves_real_overlap():
    chunks = chunk_document(
        _document((" ".join(f"w{i}" for i in range(120))).encode()),
        target_words=50,
        overlap_words=10,
    )
    assert len(chunks) == 3
    assert chunks[0].content.split()[-10:] == chunks[1].content.split()[:10]
    assert chunks[1].content.split()[-10:] == chunks[2].content.split()[:10]


def test_large_single_block_is_bounded():
    chunks = chunk_document(
        _document((" ".join(f"w{i}" for i in range(250))).encode()),
        target_words=80,
        overlap_words=10,
    )
    assert all(len(c.content.split()) <= 80 for c in chunks)
    assert len(chunks) >= 4


def test_chunk_provenance_is_present():
    chunks = chunk_document(
        _document(b"alpha beta gamma " * 100),
        target_words=40,
        overlap_words=5,
    )
    for chunk in chunks:
        assert chunk.provenance["content_hash"] == "x"
        assert chunk.provenance["block_ids"]
        assert chunk.provenance["first_block"]
        assert chunk.provenance["last_block"]


def test_invalid_chunk_parameters():
    doc = _document(b"alpha beta")
    with pytest.raises(ValueError):
        chunk_document(doc, target_words=0)
    with pytest.raises(ValueError):
        chunk_document(doc, target_words=10, overlap_words=10)


def test_document_metadata_version_foundation():
    metadata = DocumentMetadata(
        document_id="policy-001",
        title="Policy",
        created_at="2026-09-20T00:00:00+00:00",
        updated_at="2026-09-20T00:00:00+00:00",
        version=2,
        document_date="2026-09-20",
        effective_from="2026-10-01",
        project_id="project-001",
        access_scope="INTERNAL",
    )
    assert metadata.version == 2
    assert metadata.document_date == "2026-09-20"
    assert metadata.effective_from == "2026-10-01"
    assert metadata.lifecycle_status.value == "INGESTING"


def test_chunk_id_includes_version():
    doc = _document(b"alpha beta gamma")
    doc.metadata.version = 2
    chunks = chunk_document(doc)
    assert chunks[0].chunk_id == "d1:v2:chunk:00001"


def test_chunk_carries_version_metadata():
    doc = _document(b"alpha beta gamma")
    doc.metadata.version = 2
    doc.metadata.document_date = "2026-09-20"
    doc.metadata.effective_from = "2026-10-01"
    chunks = chunk_document(doc)
    assert chunks[0].document_id == "d1"
    assert chunks[0].version == 2
    assert chunks[0].document_date == "2026-09-20"
    assert chunks[0].effective_from == "2026-10-01"
    assert chunks[0].lifecycle_status.value == "INGESTING"
    assert chunks[0].content


def test_csv_extraction():
    data = "Име;Количество;Цена\nЛаптоп;2;1500\nМонитор;3;500\n".encode()
    blocks = extract_document("items.csv", data)
    assert len(blocks) == 1
    assert blocks[0].block_type == "table"
    assert "Име | Количество | Цена" in blocks[0].content
    assert "Лаптоп | 2 | 1500" in blocks[0].content
    assert blocks[0].provenance["source_format"] == "csv"


def test_docx_extraction():
    from docx import Document

    buffer = io.BytesIO()
    document = Document()
    document.add_heading("Командировки", level=1)
    document.add_paragraph("Дневните командировъчни са 40 EUR.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Разход"
    table.cell(0, 1).text = "Лимит"
    table.cell(1, 0).text = "Хотел"
    table.cell(1, 1).text = "100 EUR"
    document.save(buffer)

    blocks = extract_document("policy.docx", buffer.getvalue())
    assert len(blocks) == 3
    assert blocks[0].block_type == "heading"
    assert blocks[1].section == "Командировки"
    assert blocks[2].block_type == "table"
    assert "Разход | Лимит" in blocks[2].content
    assert blocks[2].provenance["source_format"] == "docx"


def test_xlsx_extraction():
    from openpyxl import Workbook

    buffer = io.BytesIO()
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Бюджет"
    worksheet.append(["Артикул", "Количество", "Цена"])
    worksheet.append(["Лаптоп", 2, 1500])
    worksheet.append(["Монитор", 3, 500])
    workbook.save(buffer)

    blocks = extract_document("budget.xlsx", buffer.getvalue())
    assert len(blocks) == 1
    assert blocks[0].block_type == "table"
    assert blocks[0].section == "Бюджет"
    assert "Лаптоп | 2 | 1500" in blocks[0].content
    assert blocks[0].provenance["sheet"] == "Бюджет"


def test_pptx_extraction():
    from pptx import Presentation

    buffer = io.BytesIO()
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    title = slide.shapes.title
    title.text = "Проект"
    textbox = slide.shapes.add_textbox(1000000, 1000000, 5000000, 1000000)
    textbox.text = "Бюджетът е 100 000 EUR."
    table = slide.shapes.add_table(2, 2, 1000000, 2500000, 5000000, 2000000).table
    table.cell(0, 0).text = "Разход"
    table.cell(0, 1).text = "Сума"
    table.cell(1, 0).text = "Хардуер"
    table.cell(1, 1).text = "50 000 EUR"

    presentation.save(buffer)
    blocks = extract_document("brief.pptx", buffer.getvalue())
    assert len(blocks) == 3
    assert blocks[0].block_type == "heading"
    assert blocks[1].page == 1
    assert blocks[2].block_type == "table"
    assert "Хардуер | 50 000 EUR" in blocks[2].content
    assert blocks[2].provenance["slide"] == 1


def test_pdf_extraction():
    import fitz

    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Corporate AI PDF test")
    data = pdf.tobytes()
    pdf.close()

    blocks = extract_document("test.pdf", data)

    assert len(blocks) == 1
    assert blocks[0].block_type == "text"
    assert blocks[0].page == 1
    assert blocks[0].content == "Corporate AI PDF test"
    assert blocks[0].provenance["source_format"] == "pdf"
    assert blocks[0].provenance["page"] == 1
    assert blocks[0].provenance["extraction"] == "pymupdf_text"
    assert len(blocks[0].provenance["bbox"]) == 4


def test_unsupported_extractor():
    with pytest.raises(ValueError, match="UNSUPPORTED_FILE_TYPE"):
        extract_document("image.png", b"data")


def test_pipeline_repository_knowledge_critical_path(monkeypatch):
    calls = []

    class FakeRepository:
        async def register_version(self, metadata, source_file, content_hash):
            calls.append(("register", metadata.document_id, source_file))
            metadata.version = 1
            return metadata

        async def store_canonical_source(
            self, document_id, version, filename, data, content_hash
        ):
            calls.append(("store", document_id, version, filename))
            return f"documents/{document_id}/original/{version}/{filename}"

        async def set_canonical_storage_key(
            self, document_id, version, canonical_storage_key
        ):
            calls.append(("set_key", document_id, version, canonical_storage_key))

        async def finalize_version(self, document_id, version):
            calls.append(("finalize", document_id, version))
            return DocumentMetadata(
                document_id=document_id,
                title="critical-path",
                created_at="now",
                updated_at="now",
                version=version,
                lifecycle_status=LifecycleStatus.CURRENT,
            )

        async def fail_version(self, document_id, version):
            calls.append(("fail", document_id, version))

        async def delete_canonical_source(self, document_id, version, filename):
            calls.append(("delete", document_id, version, filename))

    class FakeKnowledgeClient:
        async def ingest(self, chunk):
            calls.append(("index", chunk.document_id, chunk.version))

        async def set_lifecycle_status(self, document_id, version, lifecycle_status):
            calls.append(("lifecycle", document_id, version, lifecycle_status))
            return {"status": "ok"}

    monkeypatch.setattr(pipeline, "repository", FakeRepository())
    monkeypatch.setattr(pipeline, "KnowledgeEngineClient", FakeKnowledgeClient)

    result = asyncio.run(
        pipeline.ingest_document(
            filename="critical-path.txt",
            data=b"Repository critical path",
            document_id="critical-path-001",
        )
    )

    assert result.status.value == "READY"
    assert result.lifecycle_status.value == "CURRENT"
    assert result.chunk_count == 1
    assert result.indexed_count == 1
    assert calls == [
        ("register", "critical-path-001", "critical-path.txt"),
        ("store", "critical-path-001", 1, "critical-path.txt"),
        (
            "set_key",
            "critical-path-001",
            1,
            "documents/critical-path-001/original/1/critical-path.txt",
        ),
        ("index", "critical-path-001", 1),
        ("finalize", "critical-path-001", 1),
        ("lifecycle", "critical-path-001", 1, "CURRENT"),
    ]
